"""Compare extracted tabs against the master Word list.

Usage: python compare.py <tabs_json> <master_docx> <out_json>

Reads the first table of the master docx (columns: Tab | Title | Vehicle).
Buckets each extracted tab as:
  - matched:   already in the master (normalized-title match)
  - uncertain: similar to an existing entry but not a clear match - needs review
  - new:       not in the master, should be appended

If the master docx does not exist, every tab is "new".
"""
import sys, re, json, os
from difflib import SequenceMatcher


def norm(s):
    return re.sub(r"[^a-z0-9]", "", s.lower())


def is_match(a, b):
    if not a or not b:
        return False
    if a == b:
        return True
    if len(a) >= 40 and len(b) >= 40 and (a.startswith(b) or b.startswith(a)):
        return True
    return False


def read_master_titles(master_docx):
    """Rows already in the master list, as {tab, title, vehicle}.
    Empty when the master does not exist yet. Shared with fetch_channel.py so
    its "already indexed" pre-filter can never disagree with this comparison.

    Raises ValueError on an unusable master. It must NOT exit the process:
    fetch_channel.py calls this only as an optimisation and has to be able to
    fall back to fetching everything.
    """
    if not os.path.exists(master_docx):
        return []
    import docx
    d = docx.Document(master_docx)
    if not d.tables:
        raise ValueError("master docx has no table")
    existing = []
    for row in d.tables[0].rows[1:]:
        cells = [c.text.strip() for c in row.cells]
        if len(cells) >= 3 and cells[1] and cells[1] != "__TITLE__":
            existing.append({"tab": cells[0], "title": cells[1], "vehicle": cells[2]})
    return existing


def already_indexed(title, existing_norms):
    """True when `title` is a confident match for a master row - the same test
    main() uses for the "matched" bucket, so a caller that skips these is
    guaranteed not to turn them into "new"."""
    tn = norm(title)
    return any(is_match(tn, n_) for n_, _ in existing_norms)


def main():
    tabs_json, master_docx, out_json = sys.argv[1], sys.argv[2], sys.argv[3]
    data = json.load(open(tabs_json, encoding="utf-8"))
    tabs = data["tabs"]

    try:
        existing = read_master_titles(master_docx)
    except ValueError as e:
        print(f"ERROR: {e}", file=sys.stderr)
        sys.exit(1)
    existing_norms = [(norm(e["title"]), e) for e in existing]

    matched, uncertain, new = [], [], []
    for t in tabs:
        tn = norm(t["title"])
        hit = next((e for n_, e in existing_norms if is_match(tn, n_)), None)
        if hit:
            matched.append({"tab": t["tab"], "title": t["title"], "master_title": hit["title"]})
            continue
        best, best_ratio = None, 0.0
        for n_, e in existing_norms:
            r = SequenceMatcher(None, tn, n_).ratio()
            if r > best_ratio:
                best, best_ratio = e, r
        same_tab = next((e["title"] for e in existing if e["tab"] == str(t["tab"])), None)
        if best_ratio >= 0.72:
            uncertain.append({"tab": t["tab"], "title": t["title"], "vic": t["vic"],
                              "similar_to": best["title"], "similarity": round(best_ratio, 2),
                              "master_entry_with_same_tab_number": same_tab,
                              "snippet": t["snippet"][:600]})
        else:
            new.append({**t, "snippet": t["snippet"][:1500],
                        "master_entry_with_same_tab_number": same_tab})

    result = {"existing_count": len(existing), "matched": len(matched),
              "uncertain": uncertain, "new": new, "anomalies": data.get("anomalies", [])}
    with open(out_json, "w", encoding="utf-8") as f:
        json.dump(result, f, indent=1, ensure_ascii=False)
    print(f"existing in master: {len(existing)} | matched: {len(matched)} | "
          f"uncertain: {len(uncertain)} | new: {len(new)}")


if __name__ == "__main__":
    main()
