"""Append rows to the master Word list, preserving table formatting.

Usage: python append_master.py <master_docx> <rows_json>

rows_json is a JSON list of {"tab": <int|str>, "title": str, "vehicle": str}.

If the master docx does not exist, it is created from the bundled
assets/master_template.docx (which contains the heading, the styled header
row, and one placeholder data row used as the formatting source).

New rows are cloned from the last data row so borders, fonts, shading and
column widths carry over exactly.
"""
import sys, os, json, shutil, copy
import docx

TEMPLATE = os.path.join(os.path.dirname(__file__), "..", "assets", "master_template.docx")
PLACEHOLDER = "__TAB__"


def set_cell_text(cell, text):
    # keep the first run of the first paragraph (carries formatting), drop the rest
    para = cell.paragraphs[0]
    for extra in cell.paragraphs[1:]:
        extra._element.getparent().remove(extra._element)
    runs = para.runs
    if not runs:
        para.add_run(text)
        return
    runs[0].text = text
    for r in runs[1:]:
        r._element.getparent().remove(r._element)


def main():
    master, rows_json = sys.argv[1], sys.argv[2]
    rows = json.load(open(rows_json, encoding="utf-8"))
    if not rows:
        print("nothing to append")
        return

    created = False
    if not os.path.exists(master):
        shutil.copyfile(TEMPLATE, master)
        created = True

    d = docx.Document(master)
    table = d.tables[0]

    # if the template placeholder row is still present, fill it with the first entry
    if table.rows[-1].cells[0].text.strip() == PLACEHOLDER:
        first = rows.pop(0)
        r = table.rows[-1]
        set_cell_text(r.cells[0], str(first["tab"]))
        set_cell_text(r.cells[1], first["title"])
        set_cell_text(r.cells[2], first["vehicle"])
        appended = 1
    else:
        appended = 0

    source_tr = table.rows[-1]._element
    for entry in rows:
        new_tr = copy.deepcopy(source_tr)
        source_tr.addnext(new_tr)
        source_tr = new_tr
        from docx.table import _Row
        row = _Row(new_tr, table)
        set_cell_text(row.cells[0], str(entry["tab"]))
        set_cell_text(row.cells[1], entry["title"])
        set_cell_text(row.cells[2], entry["vehicle"])
        appended += 1

    d.save(master)
    verb = "created" if created else "updated"
    print(f"{verb} {master}: appended {appended} row(s), table now has "
          f"{len(table.rows) - 1} entries")


if __name__ == "__main__":
    main()
